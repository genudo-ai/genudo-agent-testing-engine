#!/usr/bin/env python3
"""
Test Suite Compiler for GenuDo Agents:
Validates test journey definitions and compiles them into multi-tab Excel workbooks
(Cases, Steps, GlobalRules, Facts, BrandTerms).
"""

import os
import sys
import json
import argparse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NAVY = "1B365D"
WHITE = "FFFFFF"
GRAY_BG = "F3F4F6"
BORDER_CLR = "D1D5DB"

def validate_suite(suite_data):
    """Enforces journey quality rules."""
    errors = []
    cases = suite_data.get("cases", [])
    if not cases:
        errors.append("Suite contains no test cases.")

    for c in cases:
        cid = c.get("case_id", "UNKNOWN")
        steps = c.get("steps", [])
        if len(steps) < 4:
            errors.append(f"Case {cid} has only {len(steps)} steps (minimum required: 4).")
        if len(steps) > 15:
            errors.append(f"Case {cid} has {len(steps)} steps (maximum: 15).")
        if not c.get("contact_name"):
            errors.append(f"Case {cid} is missing a contact persona name.")
        if c.get("priority") not in ["P0", "P1", "P2"]:
            errors.append(f"Case {cid} has invalid priority '{c.get('priority')}'.")

    bt = suite_data.get("brand_terms", [])
    if not isinstance(bt, list) or not all(isinstance(t, str) for t in bt):
        errors.append("brand_terms must be a list of strings.")
    return errors

def build_excel_suite(suite_data, output_path="test-suite.xlsx"):
    wb = openpyxl.Workbook()
    thin_border = Border(
        left=Side(style='thin', color=BORDER_CLR),
        right=Side(style='thin', color=BORDER_CLR),
        top=Side(style='thin', color=BORDER_CLR),
        bottom=Side(style='thin', color=BORDER_CLR)
    )

    # 1. Cases Sheet
    ws_cases = wb.active
    ws_cases.title = "Cases"
    case_headers = ["case_id", "short_name", "contact_name", "category", "priority", "pipelines", "start_stage", "description", "source_incident"]
    ws_cases.append(case_headers)

    for c in suite_data.get("cases", []):
        ws_cases.append([
            c.get("case_id", ""),
            c.get("short_name", ""),
            c.get("contact_name", ""),
            c.get("category", ""),
            c.get("priority", "P1"),
            c.get("pipelines", "all"),
            c.get("start_stage", ""),
            c.get("description", ""),
            c.get("source_incident", "")
        ])

    # 2. Steps Sheet
    ws_steps = wb.create_sheet(title="Steps")
    step_headers = [
        "case_id", "step", "customer_message", "image", "expected_reply",
        "must_contain_all", "must_contain_any", "must_not_contain",
        "regex_must", "regex_forbidden", "expect_action", "expect_payload_contains",
        "expect_stage", "expect_stage_not", "judge_rubric", "severity", "language"
    ]
    ws_steps.append(step_headers)

    for c in suite_data.get("cases", []):
        cid = c.get("case_id", "")
        for idx, s in enumerate(c.get("steps", []), 1):
            ws_steps.append([
                cid,
                s.get("step_no") or s.get("step") or idx,
                s.get("customer_message", ""),
                s.get("image", ""),
                s.get("expected_reply", ""),
                s.get("must_contain_all", ""),
                s.get("must_contain_any", ""),
                s.get("must_not_contain", ""),
                s.get("regex_must", ""),
                s.get("regex_forbidden", ""),
                s.get("expect_action", ""),
                s.get("expect_payload_contains", ""),
                s.get("expect_stage", ""),
                s.get("expect_stage_not", ""),
                s.get("judge_rubric", ""),
                s.get("severity", "major"),
                s.get("language", "")
            ])

    # 3. GlobalRules Sheet
    ws_rules = wb.create_sheet(title="GlobalRules")
    rule_headers = ["rule_id", "description", "type", "pattern", "exempt_if_reply_contains", "severity", "category"]
    ws_rules.append(rule_headers)

    for r in suite_data.get("global_rules", []):
        ws_rules.append([
            r.get("rule_id", ""),
            r.get("description", ""),
            r.get("type", ""),
            r.get("pattern", ""),
            r.get("exempt_if_reply_contains", ""),
            r.get("severity", "critical"),
            r.get("category", "")
        ])

    # 4. Facts Sheet
    ws_facts = wb.create_sheet(title="Facts")
    ws_facts.append(["Business Facts & Ground Truth"])
    for line in suite_data.get("facts", "").split("\n"):
        ws_facts.append([line])

    # 5. BrandTerms Sheet (names written identically in every language; ignored by language detection)
    ws_brand = wb.create_sheet(title="BrandTerms")
    ws_brand.append(["term"])
    for t in suite_data.get("brand_terms", []):
        ws_brand.append([t])

    # Styling all sheets
    for ws in [ws_cases, ws_steps, ws_rules, ws_facts, ws_brand]:
        for cell in ws[1]:
            cell.font = Font(name="Calibri", size=11, bold=True, color=WHITE)
            cell.fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 26

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = thin_border
                cell.font = Font(name="Calibri", size=10)

        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            max_len = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[letter].width = min(max(max_len + 3, 12), 50)

    wb.save(output_path)
    print(f"Test suite Excel workbook compiled: {output_path}")

def main():
    parser = argparse.ArgumentParser(description="GenuDo Test Suite Compiler")
    parser.add_argument("--input", default="templates/starter-suite.json", help="Input suite JSON")
    parser.add_argument("--output-xlsx", default="templates/starter-suite.xlsx", help="Output Excel file")
    parser.add_argument("--validate", action="store_true", help="Validate suite integrity only")
    args = parser.parse_args()

    if not os.path.exists(args.input):
        print(f"Input file not found: {args.input}", file=sys.stderr)
        sys.exit(1)

    with open(args.input, "r", encoding="utf-8") as f:
        data = json.load(f)

    errs = validate_suite(data)
    if errs:
        print("Suite validation failed with errors:", file=sys.stderr)
        for e in errs:
            print(f"  - {e}", file=sys.stderr)
        sys.exit(1)

    print("Suite validation PASSED.")
    if not args.validate:
        build_excel_suite(data, args.output_xlsx)

if __name__ == "__main__":
    main()
