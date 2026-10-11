#!/usr/bin/env python3
"""
GenuDo Headless Test Automation Engine:
Simulates multi-turn customer journeys against live or staging GenuDo AI agents.
Enforces is_test isolation, quiet-period polling, MCP rate limits, dual-tier evaluation,
and safe opportunity archival/pausing.
"""

import os
import re
import sys
import time
import json
import uuid
import mimetypes
import argparse
import datetime as dt
import threading
import urllib.request
import urllib.error
from concurrent.futures import ThreadPoolExecutor

if __package__ is None or __package__ == "":
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    from config import (
        GENUDO_BASE, GENUDO_TOKEN, GENUDO_ENV,
        GENUDO_ARCHIVED_STAGE, MCP_MAX_RPM,
        get_pipelines, check_prod_safety
    )
    from judge import run_checks, evaluate_with_llm, resolve_judge_backend, user_lang
    from report_generator import generate_html_report
else:
    from .config import (
        GENUDO_BASE, GENUDO_TOKEN, GENUDO_ENV,
        GENUDO_ARCHIVED_STAGE, MCP_MAX_RPM,
        get_pipelines, check_prod_safety
    )
    from .judge import run_checks, evaluate_with_llm, resolve_judge_backend, user_lang
    from .report_generator import generate_html_report

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# ------------------------------------------------------------------ Timing / Rate Limiting
# The platform allows ~30 MCP requests/min; 2.2s spacing keeps a safety margin under that.
MCP_MIN_INTERVAL = max(2.2, 60.0 / MCP_MAX_RPM)
MCP_BACKOFF_SECONDS = 25      # sleep after a -32029 / HTTP 429 rate-limit response
POLL_SECONDS = 3.5            # conversation polling interval while waiting for the AI
QUIET_SECONDS = 14            # no new message for this long = the AI finished answering
REPLY_TIMEOUT = 150           # max wait per customer message

_mcp_lock = threading.Lock()
_last_mcp_call = 0.0

def _auth_header():
    return GENUDO_TOKEN if GENUDO_TOKEN.startswith("Bearer") else f"Bearer {GENUDO_TOKEN}"

def mcp(tool_name, arguments, retries=6):
    """GenuDo MCP tool call (JSON-RPC over {GENUDO_BASE}/mcp) with a global throttle and rate-limit backoff."""
    global _last_mcp_call
    if not GENUDO_TOKEN:
        raise RuntimeError("GENUDO_TOKEN is not set.")

    body = json.dumps({
        "jsonrpc": "2.0",
        "id": 1,
        "method": "tools/call",
        "params": {"name": tool_name, "arguments": arguments}
    }).encode("utf-8")

    headers = {
        "Authorization": _auth_header(),
        "Content-Type": "application/json",
        "Accept": "application/json, text/event-stream"
    }

    last_err = None
    for attempt in range(retries):
        with _mcp_lock:
            wait = _last_mcp_call + MCP_MIN_INTERVAL - time.time()
            if wait > 0:
                time.sleep(wait)
            _last_mcp_call = time.time()

        req = urllib.request.Request(f"{GENUDO_BASE}/mcp", data=body, method="POST", headers=headers)
        try:
            with urllib.request.urlopen(req, timeout=120) as resp:
                raw = resp.read().decode("utf-8")
            data = json.loads(raw[raw.find("{"):])
        except urllib.error.HTTPError as e:
            last_err = f"HTTP {e.code}"
            time.sleep(MCP_BACKOFF_SECONDS if e.code == 429 else 5)
            continue
        except Exception as e:
            last_err = str(e)
            time.sleep(5)
            continue

        err = data.get("error")
        if err:
            if err.get("code") == -32029:  # MCP rate limit
                last_err = "rate limited (-32029)"
                time.sleep(MCP_BACKOFF_SECONDS)
                continue
            raise RuntimeError(f"MCP {tool_name} error {err.get('code')}: {err.get('message', err)}")
        res = data.get("result", {})
        return res.get("structuredContent", res)

    raise RuntimeError(f"MCP tool '{tool_name}' failed after {retries} attempts ({last_err}).")

def http(method, path, data=None, headers=None):
    """GenuDo REST call ({GENUDO_BASE}/api/user/...). Not retried, so a message is never double-sent."""
    h = {"Authorization": _auth_header(), "Accept": "application/json"}
    if headers:
        h.update(headers)
    url = f"{GENUDO_BASE}{path}" if path.startswith("/") else f"{GENUDO_BASE}/{path}"
    req = urllib.request.Request(url, data=data, method=method, headers=h)
    try:
        with urllib.request.urlopen(req, timeout=90) as resp:
            raw = resp.read().decode("utf-8")
            return json.loads(raw) if raw else {}
    except urllib.error.HTTPError as e:
        raise RuntimeError(f"{method} {path} -> HTTP {e.code}: {e.read().decode('utf-8', 'replace')[:300]}")

def multipart(fields, file_path=None):
    boundary = "----WebKitFormBoundary" + uuid.uuid4().hex[:16]
    parts = []
    for k, v in fields.items():
        parts.append(f'--{boundary}\r\nContent-Disposition: form-data; name="{k}"\r\n\r\n{v}\r\n'.encode("utf-8"))
    if file_path and os.path.exists(file_path):
        name = os.path.basename(file_path)
        mime = mimetypes.guess_type(name)[0] or "application/octet-stream"
        with open(file_path, "rb") as f:
            parts.append(
                f'--{boundary}\r\nContent-Disposition: form-data; name="files[]"; filename="{name}"\r\nContent-Type: {mime}\r\n\r\n'.encode("utf-8")
                + f.read() + b"\r\n"
            )
    parts.append(f"--{boundary}--\r\n".encode("utf-8"))
    return b"".join(parts), {"Content-Type": f"multipart/form-data; boundary={boundary}"}

def mark_test_name(msg, contact, label):
    if not msg or label in msg:
        return msg
    m = re.search(r"(الاسم\s*[:：]\s*)([^،,\n]+)", msg) or re.search(r"(اسمي\s+)([^،,\n\d]+)", msg) or re.search(r"(name\s*[:：]?\s*)([^,\n\d]+)", msg, re.I)
    if m:
        return msg[:m.end(2)].rstrip() + f" - {label}" + msg[m.end(2):]
    if contact and contact in msg:
        i = msg.index(contact) + len(contact)
        return msg[:i] + f" - {label}" + msg[i:]
    return msg

# ------------------------------------------------------------------ Suite loading
def load_suite(path):
    """Load a suite from .json or .xlsx (tabs: Cases, Steps, GlobalRules, Facts, optional BrandTerms)."""
    if path.endswith(".json"):
        with open(path, "r", encoding="utf-8") as f:
            suite = json.load(f)
        suite.setdefault("global_rules", [])
        suite.setdefault("brand_terms", [])
        return suite
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError("Reading Excel suites requires `openpyxl`. Install via `pip install openpyxl`.")
    wb = openpyxl.load_workbook(path, data_only=True)

    def read_tab(name):
        if name not in wb.sheetnames:
            return []
        ws = wb[name]
        headers = [str(c.value or "").strip() for c in ws[1]]
        return [dict(zip(headers, ["" if c.value is None else str(c.value) for c in r])) for r in ws.iter_rows(min_row=2) if any(c.value not in (None, "") for c in r)]

    cases = read_tab("Cases")
    steps = read_tab("Steps")
    for c in cases:
        c["steps"] = sorted([s for s in steps if s["case_id"] == c["case_id"]], key=lambda x: int(float(x.get("step") or 0)))
    facts = "\n".join(str(r[0].value or "") for r in wb["Facts"].iter_rows(min_row=2)) if "Facts" in wb.sheetnames else ""
    brand_terms = [r.get("term", "") for r in read_tab("BrandTerms") if r.get("term")]
    return {"cases": cases, "global_rules": read_tab("GlobalRules"), "facts": facts, "brand_terms": brand_terms}

def case_matches_channel(case, channel):
    allowed = [x.strip().lower() for x in str(case.get("pipelines") or "all").split(",")]
    return any(x in ("all", "both", "", channel) for x in allowed)

# ------------------------------------------------------------------ Engine Runner Class
class GenudoTestEngine:
    def __init__(self, suite_path, run_no=None, workers=3, stop_on_fail=False, keep_live=False, allow_prod=False, dry_run=False):
        check_prod_safety(allow_prod)
        if not dry_run and not GENUDO_TOKEN:
            sys.exit("GENUDO_TOKEN is not set (export it or put it in .env). Use --dry-run to preview without a token.")
        self.suite_path = suite_path
        self.run_no = run_no or int(time.time() % 10000)
        self.workers = workers
        self.stop_on_fail = stop_on_fail
        self.keep_live = keep_live
        self.dry_run = dry_run

        now_str = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
        self.run_id = f"R{self.run_no}-{now_str}"
        self.run_dir = os.path.join(HERE, "runs", self.run_id)

        self.suite = load_suite(suite_path)
        self.pipelines = get_pipelines()
        self.judge_backend = "off" if dry_run else resolve_judge_backend()
        self.tracked_opps = []
        self._tracked_lock = threading.Lock()
        self._stage_cache = {}
        self._stage_lock = threading.Lock()

    # -------------------------------------------------------------- platform helpers
    def stages(self, pid):
        """{stage_id: stage_name} for a pipeline (cached)."""
        with self._stage_lock:
            if pid not in self._stage_cache:
                st = mcp("list_pipeline_stages", {"pipeline_id": pid}).get("stages", [])
                self._stage_cache[pid] = {s["id"]: s["name"] for s in st}
            return self._stage_cache[pid]

    def conv_info(self, conv_id):
        """(messages_count, current opportunity stage_id) for a conversation."""
        d = http("GET", f"/api/user/conversations/{conv_id}")
        d = d.get("data", d)
        opp = (d.get("opportunities") or [{}])[0]
        return int(d.get("messages_count") or 0), opp.get("stage_id")

    def wait_for_quiet_period(self, conv_id, before):
        """Poll until the AI has answered (count >= before + 2: the customer's message plus at least one reply)
        and no new message arrived for QUIET_SECONDS. Gives up after REPLY_TIMEOUT."""
        t0 = time.time()
        last_change, last_count = time.time(), before
        while time.time() - t0 < REPLY_TIMEOUT:
            time.sleep(POLL_SECONDS)
            try:
                cnt, _ = self.conv_info(conv_id)
            except Exception:
                continue
            if cnt != last_count:
                last_count, last_change = cnt, time.time()
            if last_count >= before + 2 and time.time() - last_change >= QUIET_SECONDS:
                return True
        return last_count >= before + 2

    # -------------------------------------------------------------- one step
    def run_step(self, step, pid, conv_id, contact_id, opp_id, history, case_contact, state):
        out = {
            "step": step.get("step_no") or step.get("step"),
            "customer_message": step.get("customer_message", ""),
            "image": step.get("image", ""),
            "expected_reply": step.get("expected_reply", ""),
            "replies": [],
            "checks": [],
            "actions": [],
            "stage": None,
            "judge": None,
            "verdict": "pass"
        }

        # 1. Read initial state
        init_count, _ = self.conv_info(conv_id)
        sent_at = dt.datetime.now(dt.timezone.utc)

        # 2. Format customer text (customer-typed name gets the "TEST R<n>" suffix) and fixture
        label = f"TEST R{self.run_no}"
        msg_text = mark_test_name(step.get("customer_message") or "", case_contact, label)
        img_path = os.path.join(HERE, "templates", "fixtures", step["image"]) if step.get("image") else None
        if img_path and not os.path.exists(img_path):
            raise RuntimeError(f"fixture not found: {img_path}")

        # 3. Dispatch message
        body, h = multipart({"message": msg_text} if msg_text else {}, img_path)
        http("POST", f"/api/user/test-conversations/{conv_id}/send-message", body, h)

        # 4. Wait for quiet period
        self.wait_for_quiet_period(conv_id, init_count)

        # 5. Fetch new replies via MCP
        iso_cutoff = (sent_at - dt.timedelta(seconds=3)).strftime("%Y-%m-%dT%H:%M:%S")
        msgs = mcp("list_messages", {"pipeline_id": pid, "contact_id": contact_id, "start_date": iso_cutoff[:10], "limit": 100}).get("messages", [])
        replies = [m for m in sorted(msgs, key=lambda m: m.get("created_at", ""))
                   if m.get("role") == "assistant" and m.get("created_at", "")[:19] > iso_cutoff
                   and m.get("status") != "failed" and m.get("conversation_id", conv_id) == conv_id]
        out["replies"] = [m.get("message") or m.get("content") or "" for m in replies]

        # 6. Fetch actions. Each action result is attributed to the first step that sees it, so late webhooks
        #    are not double-counted; when an action is expected, re-poll a few times for slow webhooks.
        exp = (step.get("expect_action") or "").strip().lower()
        fired = []
        for _ in range(5 if exp and exp != "none" else 1):
            res = mcp("list_action_results", {"pipeline_id": pid, "limit": 40}).get("results", [])
            fired = [a for a in res if a.get("opportunity_id") == opp_id and a.get("id") not in state["seen_actions"]]
            if not exp or exp == "none" or any((a.get("action_name") or "").lower().startswith(exp) for a in fired):
                break
            time.sleep(6)
        state["seen_actions"] |= {a.get("id") for a in fired}
        out["actions"] = [a.get("action_name", "") for a in fired]
        out["action_payloads"] = [a.get("payload") for a in fired]

        # 7. Current stage
        _, stage_id = self.conv_info(conv_id)
        stage_name = self.stages(pid).get(stage_id) if stage_id is not None else None
        out["stage"] = stage_name

        # 8. Tier 1: deterministic checks
        raw_msg = step.get("customer_message", "")
        brand_terms = self.suite.get("brand_terms", [])
        checks = run_checks(step, out["replies"], raw_msg, self.suite.get("global_rules", []), fired, stage_name, state["lang"], brand_terms)
        out["checks"] = checks
        state["lang"] = (step.get("language") or "").lower() or user_lang(raw_msg, brand_terms) or state["lang"]

        # 9. Tier 2: LLM judge via subscription CLI (claude / codex); judge errors never change the Tier-1 verdict
        hist_txt = "\n".join(history[-12:]) or "(first message)"
        j = evaluate_with_llm(step, hist_txt, out["replies"], self.suite.get("facts", ""), self.judge_backend)
        out["judge"] = j
        history.append("customer: " + (raw_msg or "(image)"))
        history.extend("agent: " + t for t in out["replies"])

        sev = (step.get("severity") or "major").lower()
        failed = [c for c in checks if not c["ok"] and c["severity"] in ("critical", "major")]
        warn = [c for c in checks if not c["ok"] and c["severity"] == "minor"]
        if failed or (j and j.get("verdict") == "fail" and sev != "minor"):
            out["verdict"] = "fail"
        elif warn or (j and j.get("verdict") == "fail"):
            out["verdict"] = "warn"
        if j and j.get("verdict") == "judge_error":
            out["judge_error"] = j.get("reason", "")
        return out

    # -------------------------------------------------------------- one case
    def run_case(self, case, channel):
        pid = self.pipelines.get(channel)
        res = {
            "case_id": case["case_id"],
            "short_name": case.get("short_name", ""),
            "category": case.get("category", ""),
            "priority": case.get("priority", "P1"),
            "pipeline": channel,
            "description": case.get("description", ""),
            "steps": [],
            "verdict": "pass"
        }

        if self.dry_run:
            target = f"pipeline {pid}" if pid else "NO pipeline id in GENUDO_PIPELINES"
            print(f"  [DRY-RUN] {channel:9} {case['case_id']} {case.get('short_name')} ({len(case.get('steps', []))} steps, {target})")
            res["verdict"] = "dry-run"
            return res
        if not pid:
            res.update(verdict="error", error=f"No pipeline ID for channel '{channel}' in GENUDO_PIPELINES")
            return res

        try:
            # Create an isolated TEST opportunity (is_test, @example.com email, "TEST R<n>" name suffix)
            contact_name = f"{case.get('contact_name') or 'Customer'} - TEST R{self.run_no}"
            email = f"r{self.run_no}{case['case_id']}{channel[0]}-{uuid.uuid4().hex[:4]}@example.com".lower()
            opp_body = {
                "pipeline_id": pid,
                "name": f"Run {self.run_no} - {case['case_id']} - {case.get('short_name')}",
                "is_test": True,
                "tags": ["autotest", self.run_id],
                "notes": f"Automated test run {self.run_id}. Safe to delete.",
                "contact": {"name": contact_name, "email": email}
            }
            if case.get("start_stage"):
                rev = {v.lower(): k for k, v in self.stages(pid).items()}
                if case["start_stage"].lower() in rev:
                    opp_body["stage_id"] = rev[case["start_stage"].lower()]
            opp_data = http("POST", "/api/user/opportunities", json.dumps(opp_body).encode("utf-8"), {"Content-Type": "application/json"}).get("data", {})
            opp_id = opp_data.get("id")
            conv_id = (opp_data.get("conversation") or {}).get("id")
            contact_id = opp_data.get("contact_id")
            if opp_id:
                with self._tracked_lock:
                    self.tracked_opps.append({"opp_id": opp_id, "conv_id": conv_id, "pipeline_id": pid, "channel": channel})
            if not (opp_id and conv_id and contact_id):
                raise RuntimeError(f"unexpected create-opportunity response: {json.dumps(opp_data)[:200]}")
            res.update(opportunity_id=opp_id, conversation_id=conv_id)

            history = []
            state = {"seen_actions": set(), "lang": None}
            for st in case.get("steps", []):
                step_res = self.run_step(st, pid, conv_id, contact_id, opp_id, history, case.get("contact_name"), state)
                res["steps"].append(step_res)
                if step_res["verdict"] == "fail":
                    res["verdict"] = "fail"
                    if self.stop_on_fail:
                        break

            print(f"  [{res['verdict'].upper():5}] {channel:9} {case['case_id']} - {case.get('short_name')}", flush=True)
        except Exception as e:
            res["verdict"] = "error"
            res["error"] = str(e)
            print(f"  [ERROR] {channel:9} {case['case_id']} - Exception: {e}", file=sys.stderr, flush=True)

        return res

    # -------------------------------------------------------------- cleanup
    def cleanup(self):
        """Move every test opportunity of this run to the Archived stage and pause its conversation."""
        if self.keep_live or self.dry_run or not self.tracked_opps:
            return
        print(f"\nCleaning up {len(self.tracked_opps)} test opportunities (move to '{GENUDO_ARCHIVED_STAGE}' + pause)...")
        by_pid = {}
        for item in self.tracked_opps:
            by_pid.setdefault(item["pipeline_id"], []).append(item)
        for pid, items in by_pid.items():
            try:
                sid = next((k for k, v in self.stages(pid).items() if v.strip().lower() == GENUDO_ARCHIVED_STAGE.strip().lower()), None)
                if sid is None:
                    print(f"  WARNING: pipeline {pid} has no '{GENUDO_ARCHIVED_STAGE}' stage - opportunities NOT archived (create it).", file=sys.stderr)
                else:
                    ids = [i["opp_id"] for i in items]
                    for k in range(0, len(ids), 50):
                        mcp("update_opportunities", {"opportunity_ids": ids[k:k + 50], "stage_id": sid})
                    print(f"  archived {len(ids)} opportunities in pipeline {pid} (stage {sid})")
            except Exception as e:
                print(f"  Cleanup warning (archive, pipeline {pid}): {e}", file=sys.stderr)
            paused = 0
            for item in items:
                if not item.get("conv_id"):
                    continue
                try:
                    http("PUT", f"/api/user/conversations/{item['conv_id']}", json.dumps({"status": "paused"}).encode("utf-8"), {"Content-Type": "application/json"})
                    paused += 1
                except Exception as e:
                    print(f"  Cleanup warning (pause conv {item['conv_id']}): {e}", file=sys.stderr)
            print(f"  paused {paused} conversations in pipeline {pid}")

    # -------------------------------------------------------------- run
    def execute(self, channel="messenger", cases_filter=None, priority_filter=None, category_filter=None):
        channels = list(self.pipelines) if channel in ("all", "both") else [channel.lower()]
        print(f"\n=======================================================")
        print(f"GenuDo Test Simulation Run: {self.run_id}{'  [DRY-RUN]' if self.dry_run else ''}")
        print(f"Target Base: {GENUDO_BASE} | Env: {GENUDO_ENV} | Channels: {', '.join(channels) or 'none'} | Judge: {self.judge_backend}")
        print(f"=======================================================\n")

        priorities = [p.strip() for p in priority_filter.split(",")] if priority_filter and priority_filter != "all" else None
        jobs = []
        for ch in channels:
            for c in self.suite.get("cases", []):
                if not case_matches_channel(c, ch):
                    continue
                if cases_filter and c["case_id"] not in cases_filter:
                    continue
                if priorities and c.get("priority") not in priorities:
                    continue
                if category_filter and category_filter.lower() not in (c.get("category") or "").lower():
                    continue
                jobs.append((c, ch))

        print(f"Selected {len(jobs)} journey run(s), {sum(len(c.get('steps', [])) for c, _ in jobs)} customer messages.\n")
        results = []

        try:
            with ThreadPoolExecutor(max_workers=max(1, self.workers)) as executor:
                futures = [executor.submit(self.run_case, c, ch) for c, ch in jobs]
                for f in futures:
                    results.append(f.result())
        finally:
            self.cleanup()

        if self.dry_run:
            print("\nDry run: nothing was sent and no report was written.")
            return 0

        run_data = {
            "run_id": self.run_id,
            "timestamp": dt.datetime.now().isoformat(timespec="seconds"),
            "base_url": GENUDO_BASE,
            "suite": os.path.basename(self.suite_path),
            "judge": self.judge_backend,
            "dry_run": self.dry_run,
            "cases": results
        }

        os.makedirs(self.run_dir, exist_ok=True)
        json_path = os.path.join(self.run_dir, "results.json")
        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(run_data, f, ensure_ascii=False, indent=2)

        html_path = os.path.join(self.run_dir, "report.html")
        generate_html_report(run_data, html_path)

        passed = sum(c.get("verdict") == "pass" for c in results)
        p0_fails = [c for c in results if c.get("priority") == "P0" and c.get("verdict") != "pass"]
        print(f"\n{passed}/{len(results)} journey runs passed - P0 failures/errors: {len(p0_fails)}")
        print(f"results: {json_path}")
        return 1 if p0_fails else 0

# ------------------------------------------------------------------ compare
def compare_runs(old_path, new_path):
    def verdicts(path):
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        return {(r["case_id"], r.get("pipeline", "")): r.get("verdict") for r in data.get("cases", data.get("results", []))}
    old, new = verdicts(old_path), verdicts(new_path)
    reg = sorted(k for k in new if old.get(k) == "pass" and new[k] != "pass")
    fix = sorted(k for k in new if old.get(k) not in (None, "pass") and new[k] == "pass")
    print("REGRESSIONS (pass -> not pass):", ", ".join(f"{c}/{p}" for c, p in reg) or "none")
    print("FIXED (not pass -> pass):", ", ".join(f"{c}/{p}" for c, p in fix) or "none")
    print(f"old {sum(v == 'pass' for v in old.values())}/{len(old)}  new {sum(v == 'pass' for v in new.values())}/{len(new)}")
    return 1 if reg else 0

def main():
    for stream in (sys.stdout, sys.stderr):  # Windows consoles default to cp1252; suites are Arabic
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser(description="GenuDo Agent Test Simulation Engine")
    subparsers = parser.add_subparsers(dest="command")

    # run command
    run_p = subparsers.add_parser("run", help="Execute an automated test run")
    run_p.add_argument("--suite", default="templates/starter-suite.json", help="Path to suite file (.json or .xlsx)")
    run_p.add_argument("--pipeline", default="messenger", help="Channel name from GENUDO_PIPELINES, or 'all'")
    run_p.add_argument("--run-no", type=int, help="Run number used in 'TEST R<n>' labels, tags and the run id")
    run_p.add_argument("--cases", help="Comma-separated case IDs to run (e.g. J01,J02)")
    run_p.add_argument("--priority", help="Priority filter: P0, P1, P2, a comma list (P0,P1) or 'all'")
    run_p.add_argument("--category", help="Only cases whose category contains this text (case-insensitive)")
    run_p.add_argument("--workers", type=int, default=3, help="Concurrent cases (default: 3)")
    run_p.add_argument("--stop-on-fail", action="store_true", help="Stop a case at its first failing step")
    run_p.add_argument("--keep-live", action="store_true", help="Skip cleanup (do not archive/pause test opportunities)")
    run_p.add_argument("--allow-prod", action="store_true", help="Required when GENUDO_ENV=prod")
    run_p.add_argument("--dry-run", action="store_true", help="Print the selected journeys; no network calls")

    # list command
    list_p = subparsers.add_parser("list", help="List the test cases in a suite")
    list_p.add_argument("--suite", default="templates/starter-suite.json", help="Path to suite file (.json or .xlsx)")

    # compare command
    cmp_p = subparsers.add_parser("compare", help="Compare two runs' results.json (exit 1 on regressions)")
    cmp_p.add_argument("old", help="Older runs/<id>/results.json")
    cmp_p.add_argument("new", help="Newer runs/<id>/results.json")

    args = parser.parse_args()

    def resolve(path):
        return path if os.path.isabs(path) or os.path.exists(path) else os.path.join(HERE, path)

    if args.command == "list":
        suite_path = resolve(args.suite)
        suite = load_suite(suite_path)
        print(f"\nTest journeys in {args.suite}:")
        for c in suite.get("cases", []):
            print(f"  {c.get('case_id'):4} [{c.get('priority', 'P1'):2}] {str(c.get('pipelines') or 'all'):9} {str(c.get('category', '')):16} "
                  f"{str(c.get('short_name', '')):28} ({len(c.get('steps', []))} steps)")
        print(f"{len(suite.get('cases', []))} cases")
        sys.exit(0)

    elif args.command == "run":
        engine = GenudoTestEngine(
            suite_path=resolve(args.suite),
            run_no=args.run_no,
            workers=args.workers,
            stop_on_fail=args.stop_on_fail,
            keep_live=args.keep_live,
            allow_prod=args.allow_prod,
            dry_run=args.dry_run
        )
        cases_filter = [x.strip() for x in args.cases.split(",")] if args.cases else None
        sys.exit(engine.execute(channel=args.pipeline, cases_filter=cases_filter, priority_filter=args.priority, category_filter=args.category))

    elif args.command == "compare":
        sys.exit(compare_runs(args.old, args.new))

    else:
        parser.print_help()

if __name__ == "__main__":
    main()
